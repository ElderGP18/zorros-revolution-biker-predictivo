"""Convierte los capitulos en Markdown a un unico .docx con el formato de la tesis.

Formato observado en el PDF de la tesis: Times New Roman 12 pt en el cuerpo,
9 pt en tablas, titulos en negrita, tamano carta, captions de tablas y figuras
DEBAJO del elemento con la forma "Tabla N. Titulo. Fuente: ...".

Uso (desde backend/, con el venv activo):
    python ../docs/tesis/md_a_docx.py salida.docx cap4.md cap5.md cap6.md referencias.md
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

FUENTE = "Times New Roman"
NEGRO = RGBColor(0, 0, 0)
RE_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
RE_CAPTION = re.compile(r"^(Tabla|Figura) \d+\.")
RE_IMG = re.compile(r"^!\[(.*?)\]\((.*?)\)$")


def _fuente(run, tam=12, negrita=False, cursiva=False, mono=False):
    run.font.name = "Courier New" if mono else FUENTE
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    run.font.size = Pt(tam)
    run.font.bold = negrita
    run.font.italic = cursiva
    run.font.color.rgb = NEGRO


def _runs(parrafo, texto, tam=12, negrita=False, cursiva=False):
    """Divide el texto en runs respetando **negrita**, *cursiva* y `codigo`."""
    for trozo in RE_INLINE.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**"):
            _fuente(parrafo.add_run(trozo[2:-2]), tam, True, cursiva)
        elif trozo.startswith("`"):
            _fuente(parrafo.add_run(trozo[1:-1]), tam - 1, negrita, cursiva, mono=True)
        elif trozo.startswith("*"):
            _fuente(parrafo.add_run(trozo[1:-1]), tam, negrita, True)
        else:
            _fuente(parrafo.add_run(trozo), tam, negrita, cursiva)


def _parrafo(doc, texto, tam=12, negrita=False, cursiva=False, sangria=True, justificado=True,
             espacio_despues=6, interlineado=1.5, alineacion=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(espacio_despues)
    pf.line_spacing = interlineado
    if sangria:
        pf.first_line_indent = Cm(1.25)
    p.alignment = alineacion if alineacion is not None else (WD_ALIGN_PARAGRAPH.JUSTIFY if justificado else WD_ALIGN_PARAGRAPH.LEFT)
    _runs(p, texto, tam, negrita, cursiva)
    return p


def _titulo(doc, texto, nivel):
    tam = {1: 14, 2: 12, 3: 12}[nivel]
    p = doc.add_heading(level=nivel)
    p.paragraph_format.space_before = Pt(18 if nivel == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(texto)
    _fuente(run, tam, negrita=True, cursiva=(nivel == 3))
    return p


def _bordes(tabla):
    tbl = tabla._tbl
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl.tblPr.append(borders)


def _tabla(doc, filas):
    filas = [[c.strip() for c in f.strip().strip("|").split("|")] for f in filas]
    filas = [f for f in filas if not all(re.fullmatch(r":?-{2,}:?", c) for c in f)]
    ncol = max(len(f) for f in filas)
    t = doc.add_table(rows=len(filas), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _bordes(t)
    for i, fila in enumerate(filas):
        for j in range(ncol):
            celda = t.cell(i, j)
            celda.text = ""
            p = celda.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            _runs(p, fila[j] if j < len(fila) else "", tam=9, negrita=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def _codigo(doc, lineas, lenguaje):
    if lenguaje == "mermaid":
        _parrafo(doc, "[Diagrama en notación Mermaid: renderizar en https://mermaid.live y sustituir este bloque por la imagen.]",
                 tam=9, cursiva=True, sangria=False, interlineado=1.0)
    for ln in lineas:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Cm(0.75)
        _fuente(p.add_run(ln if ln else " "), 8, mono=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def convertir(md_paths, salida, base_figuras):
    doc = Document()
    for s in doc.sections:
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.54)
    normal = doc.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(12)

    primero = True
    for ruta in md_paths:
        lineas = Path(ruta).read_text(encoding="utf-8").splitlines()
        i, en_tabla, tabla, en_codigo, codigo, lenguaje = 0, False, [], False, [], ""
        while i < len(lineas):
            ln = lineas[i]
            if en_codigo:
                if ln.startswith("```"):
                    _codigo(doc, codigo, lenguaje); en_codigo, codigo = False, []
                else:
                    codigo.append(ln)
                i += 1; continue
            if ln.startswith("```"):
                en_codigo, lenguaje = True, ln[3:].strip(); i += 1; continue
            if ln.startswith("|"):
                tabla.append(ln); en_tabla = True; i += 1; continue
            if en_tabla:
                _tabla(doc, tabla); tabla, en_tabla = [], False
            if ln.startswith("# "):
                if not primero:
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                primero = False
                _titulo(doc, ln[2:].strip(), 1)
            elif ln.startswith("## "):
                _titulo(doc, ln[3:].strip(), 2)
            elif ln.startswith("### "):
                _titulo(doc, ln[4:].strip(), 3)
            elif RE_IMG.match(ln):
                alt, src = RE_IMG.match(ln).groups()
                archivo = Path(base_figuras) / src
                if archivo.exists():
                    doc.add_picture(str(archivo), width=Inches(6.0))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    _parrafo(doc, f"[Figura pendiente: {src}]", tam=9, cursiva=True, sangria=False)
            elif RE_CAPTION.match(ln):
                _parrafo(doc, ln, tam=10, sangria=False, justificado=False, interlineado=1.0, espacio_despues=12)
            elif ln.startswith("- "):
                p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.line_spacing = 1.5
                _runs(p, ln[2:]); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            elif re.match(r"^\d+\. ", ln):
                p = doc.add_paragraph(style="List Number"); p.paragraph_format.line_spacing = 1.5
                _runs(p, re.sub(r"^\d+\. ", "", ln)); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            elif ln.startswith("> "):
                _parrafo(doc, ln[2:], tam=11, cursiva=True, sangria=False)
            elif ln.strip():
                _parrafo(doc, ln)
            i += 1
        if en_tabla:
            _tabla(doc, tabla)
    doc.save(salida)
    return salida


if __name__ == "__main__":
    salida, *entradas = sys.argv[1:]
    figuras = Path(entradas[0]).parent / "figuras"
    print("escrito:", convertir(entradas, salida, figuras))
