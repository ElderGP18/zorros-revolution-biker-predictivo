"""Integra los Capitulos IV, V y VI al Word de la tesis usando sus propios estilos.

- Inserta los capitulos antes del titulo "Referencias".
- Los pies de tabla y figura se crean como campos SEQ, igual que los existentes,
  para que Word los numere y los incluya en los indices al actualizar campos.
- Fusiona las referencias nuevas con las existentes en orden alfabetico,
  conservando intactas las originales.
- Actualiza los dos parrafos de la Introduccion que describen el alcance.

Uso (desde backend/, con el venv):
    python ../docs/tesis/integrar_docx.py
"""
from __future__ import annotations

import copy
import re
import shutil
import sys
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips

RAIZ = Path(__file__).resolve().parents[2]
DOCS = RAIZ / "docs"
TESIS = DOCS / "tesis"
FIGURAS = TESIS / "figuras"
SCRATCH = Path(r"C:/Users/gape2/AppData/Local/Temp/claude/c--Users-gape2-Documents-zorros-revolution-biker-predictivo/b6497813-e1d2-44a1-9781-5aae7c1d575f/scratchpad")
MERMAID = SCRATCH / "mermaid"
CAPTURAS = SCRATCH / "capturas"

ACTUAL = DOCS / "Trabajo de Graduacion.docx"
SALIDA = DOCS / "Trabajo de Graduacion - integrado.docx"
RESPALDO = DOCS / "respaldo"
# Siempre se parte del original de agosto: el Word "actual" ya contiene los
# capítulos tras la primera corrida y volver a leerlo los duplicaría.
ORIGINAL = RESPALDO / "Trabajo de Graduacion - v1 agosto 2026.docx"

CAPITULOS = ["capitulo-4-resultados.md", "capitulo-5-factibilidad.md", "capitulo-6-desarrollo.md"]
ANCHO_TABLA = 9360  # twips, igual que las tablas existentes (6.5")
ANCHO_IMAGEN = Inches(6.5)
ALTO_MAX_IMAGEN = Inches(7.2)
FILAS_TABLA_CORTA = 8  # hasta este tamaño la tabla no se parte entre páginas
TAMANO_TABLA_PT = 12
FUENTE = "Times New Roman"
SOMBREADO_ENCABEZADO = "D9EAF7"  # el mismo azul claro de las tablas originales
# Ancho medio de un carácter de Times New Roman 12 más los márgenes de celda,
# en twips; con esto ninguna palabra ni cifra se parte dentro de una columna.
TWIPS_POR_CARACTER = 120
MARGEN_CELDA = 300
MAX_DEMANDA = 70

RE_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|<sub>[^<]+</sub>)")
RE_IMG = re.compile(r"^!\[(.*?)\]\((.*?)\)$")
RE_CAPTION = re.compile(r"^(Tabla|Figura) (\d+)\.(.*)$")

# Figuras 11-14 son los diagramas Mermaid (en orden de aparicion) y 15-20 las capturas.
PNG_MERMAID = ["fig11_casos_de_uso.png", "fig12_arquitectura.png", "fig13_mer.png", "fig14_recalibracion.png"]
PNG_CAPTURAS = {15: "fig15_acceso.png", 16: "fig16_centro_control.png", 17: "fig17_ventas.png",
                18: "fig18_productos.png", 19: "fig19_pronostico.png", 20: "fig20_usuarios.png"}


# ---------------------------------------------------------------------------
# Utilidades de bajo nivel
# ---------------------------------------------------------------------------

class Insertor:
    """Crea elementos al final del documento y los mueve antes del ancla."""

    def __init__(self, doc, ancla_p):
        self.doc = doc
        self.ancla = ancla_p

    def _mover(self, elemento):
        self.ancla.addprevious(elemento)

    def parrafo(self, estilo="Normal"):
        p = self.doc.add_paragraph(style=estilo)
        self._mover(p._p)
        return p

    def tabla(self, filas, cols):
        t = self.doc.add_table(rows=filas, cols=cols, style="Table Grid")
        self._mover(t._tbl)
        return t


def _runs_inline(p, texto, tam=None, negrita=False, cursiva=False):
    for trozo in RE_INLINE.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**"):
            r = p.add_run(trozo[2:-2]); r.bold = True; r.italic = cursiva
        elif trozo.startswith("`"):
            # En el cuerpo el código va en monoespaciada; en tablas, en la fuente del documento.
            r = p.add_run(trozo[1:-1]); r.font.name = FUENTE if tam else "Courier New"; r.bold = negrita
        elif trozo.startswith("<sub>"):
            r = p.add_run(trozo[5:-6]); r.font.subscript = True; r.bold = negrita
        elif trozo.startswith("*"):
            r = p.add_run(trozo[1:-1]); r.italic = True; r.bold = negrita
        else:
            r = p.add_run(trozo); r.bold = negrita; r.italic = cursiva
        if tam:
            r.font.size = Pt(tam)
            if r.font.name != "Courier New":
                r.font.name = FUENTE


def _campo_seq(p, etiqueta, numero_visible):
    """Agrega los runs de un campo SEQ tal como los escribe Word."""
    def run_fld(tipo):
        r = OxmlElement("w:r"); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), tipo); r.append(f); p._p.append(r)
    run_fld("begin")
    r = OxmlElement("w:r"); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = f" SEQ {etiqueta} \\* ARABIC "; r.append(it); p._p.append(r)
    run_fld("separate")
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr"); rpr.append(OxmlElement("w:noProof")); r.append(rpr)
    t = OxmlElement("w:t"); t.text = str(numero_visible); r.append(t); p._p.append(r)
    run_fld("end")


def caption(ins, etiqueta, numero, resto):
    """Pie 'Tabla N. ...' o 'Figura N. ...' con campo SEQ, como los existentes."""
    p = ins.parrafo("Caption")
    r = p.add_run(f"{etiqueta} "); r.bold = True
    _campo_seq(p, etiqueta, numero)
    _runs_inline(p, resto)
    return p


def _sin_sangria(p):
    p.paragraph_format.first_line_indent = 0


def _formato_celda(p):
    """Como en las tablas originales: sin sangría y a espacio sencillo."""
    _sin_sangria(p)
    p.paragraph_format.line_spacing = 1.0


def _fijar_ancho(tabla, anchos):
    tbl = tabla._tbl
    tblPr = tbl.tblPr
    w = tblPr.find(qn("w:tblW"))
    if w is None:
        w = OxmlElement("w:tblW"); tblPr.append(w)
    w.set(qn("w:w"), str(ANCHO_TABLA)); w.set(qn("w:type"), "dxa")
    lay = tblPr.find(qn("w:tblLayout"))
    if lay is None:
        lay = OxmlElement("w:tblLayout"); tblPr.append(lay)
    lay.set(qn("w:type"), "fixed")
    grid = tbl.find(qn("w:tblGrid"))
    for gc, a in zip(grid.findall(qn("w:gridCol")), anchos):
        gc.set(qn("w:w"), str(a))
    for fila in tabla.rows:
        for celda, a in zip(fila.cells, anchos):
            celda.width = Twips(a)


def _texto_plano(texto):
    return re.sub(r"\*\*|`|\*|</?sub>", "", texto)


def anchos_columnas(filas):
    """Cada columna recibe al menos lo que ocupa su palabra más larga; el
    ancho sobrante se reparte según la longitud del contenido."""
    ncol = len(filas[0])
    minimos, demanda = [], []
    for j in range(ncol):
        textos = [_texto_plano(f[j]) for f in filas]
        palabras = [len(w) for t in textos for w in t.split()] or [1]
        encabezado = max((len(w) for w in textos[0].split()), default=1) * 1.1
        minimos.append(int(max(max(palabras), encabezado) * TWIPS_POR_CARACTER) + MARGEN_CELDA)
        demanda.append(max(min(len(t), MAX_DEMANDA) for t in textos))
    total_min = sum(minimos)
    if total_min >= ANCHO_TABLA:
        anchos = [int(ANCHO_TABLA * m / total_min) for m in minimos]
    else:
        sobrante = ANCHO_TABLA - total_min
        total_dem = sum(demanda) or 1
        anchos = [m + int(sobrante * d / total_dem) for m, d in zip(minimos, demanda)]
    anchos[anchos.index(max(anchos))] += ANCHO_TABLA - sum(anchos)
    return anchos


def _sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def tabla_md(ins, filas_md):
    filas = [[c.strip() for c in f.strip().strip("|").split("|")] for f in filas_md]
    filas = [f for f in filas if not all(re.fullmatch(r":?-{2,}:?", c) for c in f)]
    ncol = max(len(f) for f in filas)
    filas = [f + [""] * (ncol - len(f)) for f in filas]
    anchos = anchos_columnas(filas)

    t = ins.tabla(len(filas), ncol)
    _fijar_ancho(t, anchos)
    es_corta = len(filas) <= FILAS_TABLA_CORTA
    for i, fila in enumerate(filas):
        if i == 0:
            trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); trPr.append(h)
        es_ultima = i == len(filas) - 1
        for j, texto in enumerate(fila):
            celda = t.cell(i, j)
            if i == 0:
                _sombrear(celda, SOMBREADO_ENCABEZADO)
            p = celda.paragraphs[0]
            _formato_celda(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            _runs_inline(p, texto, tam=TAMANO_TABLA_PT, negrita=(i == 0))
            # Una tabla corta no se parte; la última fila siempre arrastra al pie.
            if es_corta or es_ultima:
                p.paragraph_format.keep_with_next = True
    return t


def imagen(ins, ruta: Path, ancho=ANCHO_IMAGEN):
    from PIL import Image as PILImage
    with PILImage.open(ruta) as im:
        w, h = im.size
    alto = ancho * h / w
    if alto > ALTO_MAX_IMAGEN:
        ancho = ALTO_MAX_IMAGEN * w / h
    p = ins.parrafo("Normal")
    _sin_sangria(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ruta), width=int(ancho))
    return p


def caja_pendiente(ins, texto):
    t = ins.tabla(1, 1)
    _fijar_ancho(t, [ANCHO_TABLA])
    trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:trHeight"); h.set(qn("w:val"), "2600"); trPr.append(h)
    p = t.cell(0, 0).paragraphs[0]
    _formato_celda(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    r = p.add_run(texto); r.italic = True; r.font.size = Pt(TAMANO_TABLA_PT)


def _siguiente_no_vacia(lineas, i):
    while i < len(lineas) and not lineas[i].strip():
        i += 1
    return lineas[i] if i < len(lineas) else ""


def salto_de_pagina(ins):
    p = ins.parrafo("Normal")
    _sin_sangria(p)
    p.add_run().add_break(WD_BREAK.PAGE)


def parrafo_lista(ins, texto, marca):
    p = ins.parrafo("Normal")
    p.paragraph_format.left_indent = Twips(720)
    p.paragraph_format.first_line_indent = Twips(-360)
    p.add_run(marca + "\t")
    _runs_inline(p, texto)


# ---------------------------------------------------------------------------
# Conversion de un capitulo Markdown
# ---------------------------------------------------------------------------

def insertar_capitulo(ins, ruta_md: Path, estado: dict, heading1_rompe_pagina: bool):
    lineas = ruta_md.read_text(encoding="utf-8").splitlines()
    i, tabla, codigo, en_codigo, lenguaje = 0, [], [], False, ""
    while i < len(lineas):
        ln = lineas[i]
        if en_codigo:
            if ln.startswith("```"):
                if lenguaje == "mermaid":
                    idx = estado["mermaid"]; estado["mermaid"] += 1
                    png = MERMAID / PNG_MERMAID[idx] if idx < len(PNG_MERMAID) else None
                    if png and png.exists():
                        imagen(ins, png)
                    else:
                        p = ins.parrafo("Normal"); _sin_sangria(p)
                        r = p.add_run("\n".join(codigo)); r.font.name = "Courier New"; r.font.size = Pt(8)
                en_codigo, codigo = False, []
            else:
                codigo.append(ln)
            i += 1; continue
        if ln.startswith("```"):
            en_codigo, lenguaje = True, ln[3:].strip(); i += 1; continue
        if ln.startswith("|"):
            tabla.append(ln); i += 1; continue
        if tabla:
            tabla_md(ins, tabla); tabla = []
            # El pie va pegado a la tabla, como en el documento original; si no
            # hay pie, se deja la separación en blanco.
            if not RE_CAPTION.match(_siguiente_no_vacia(lineas, i)):
                ins.parrafo("Normal")

        if ln.startswith("# "):
            if not heading1_rompe_pagina:
                salto_de_pagina(ins)
            ins.parrafo("Heading 1").add_run(ln[2:].strip())
        elif ln.startswith("## "):
            ins.parrafo("Heading 2").add_run(ln[3:].strip())
        elif ln.startswith("### "):
            ins.parrafo("Heading 3").add_run(ln[4:].strip())
        elif RE_IMG.match(ln):
            _, src = RE_IMG.match(ln).groups()
            ruta = FIGURAS / src
            if ruta.exists():
                imagen(ins, ruta)
            else:
                caja_pendiente(ins, f"Figura pendiente: {src}")
        elif RE_CAPTION.match(ln):
            etiqueta, numero, resto = RE_CAPTION.match(ln).groups()
            numero = int(numero)
            if "[Insertar captura]" in resto:
                png = CAPTURAS / PNG_CAPTURAS.get(numero, "")
                if png.exists():
                    imagen(ins, png)
                else:
                    caja_pendiente(ins, "Captura de pantalla pendiente de inserción")
                resto = resto.replace(" [Insertar captura]", "").replace("[Insertar captura]", "").rstrip()
            caption(ins, etiqueta, numero, "." + resto)
        elif ln.startswith("- "):
            parrafo_lista(ins, ln[2:], "•")
        elif re.match(r"^\d+\. ", ln):
            n = re.match(r"^(\d+)\. ", ln).group(1)
            parrafo_lista(ins, re.sub(r"^\d+\. ", "", ln), f"{n}.")
        elif ln.strip():
            p = ins.parrafo("Normal")
            _runs_inline(p, ln)
        i += 1
    if tabla:
        tabla_md(ins, tabla)


# ---------------------------------------------------------------------------
# Referencias e Introduccion
# ---------------------------------------------------------------------------

def _clave(texto):
    t = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9 ]", "", t)


def fusionar_referencias(doc, heading_ref):
    nuevas = [b.strip() for b in (TESIS / "referencias-nuevas.md").read_text(encoding="utf-8").split("\n\n")
              if b.strip() and not b.startswith("#") and not b.startswith("Las referencias siguientes")]
    parrafos = doc.paragraphs
    idx = [p._p for p in parrafos].index(heading_ref._p)
    existentes = [p for p in parrafos[idx + 1:] if p.text.strip()]
    plantilla = existentes[0]

    def crear_antes(p_ref, texto):
        nuevo = p_ref.insert_paragraph_before("", style="Normal")
        if plantilla._p.pPr is not None:
            nuevo._p.insert(0, copy.deepcopy(plantilla._p.pPr))
        _runs_inline(nuevo, texto)
        return nuevo

    insertadas = 0
    for entrada in sorted(nuevas, key=_clave):
        k = _clave(entrada)
        destino = next((p for p in existentes if _clave(p.text) > k), None)
        if destino is not None:
            crear_antes(destino, entrada)
        else:
            ultimo = existentes[-1]
            nuevo = doc.add_paragraph("", style="Normal")
            if plantilla._p.pPr is not None:
                nuevo._p.insert(0, copy.deepcopy(plantilla._p.pPr))
            _runs_inline(nuevo, entrada)
            ultimo._p.addnext(nuevo._p)
            existentes.append(nuevo)
        insertadas += 1
    total = len([p for p in doc.paragraphs[idx + 1:] if p.text.strip()])
    return insertadas, total


def integrar_entregable(doc):
    """El Entregable de Ingeniería era una sección independiente con apartados
    4.x que se confundían con el Capítulo IV. Pasa a ser la sección 3.5 del
    Capítulo III: su título baja a Heading 2 y sus apartados a Heading 3 y 4,
    renumerados 3.5.x, sin tocar el contenido ni el orden de tablas y figuras."""
    dentro, cambiados = False, 0
    for p in doc.paragraphs:
        if p.style.name == "Heading 1":
            dentro = p.text.startswith("Entregable de Ingeniería")
            if dentro:
                p.style = doc.styles["Heading 2"]
                p.paragraph_format.page_break_before = None
                p.paragraph_format.space_after = None
                p.runs[0].text = "3.5 " + p.runs[0].text
                cambiados += 1
            continue
        if not (dentro and p.runs and p.runs[0].text.startswith("4.")):
            continue
        if p.style.name == "Heading 2":
            p.style = doc.styles["Heading 3"]
        elif p.style.name == "Heading 3":
            p.style = doc.styles["Heading 4"]
        else:
            continue
        p.runs[0].text = "3.5." + p.runs[0].text[2:]
        cambiados += 1
    return cambiados


def ajustar_estilos(doc):
    """Subtítulos x.y.z con una tabulación (y dos en el cuarto nivel); pies de
    tabla y figura en la misma fuente y tamaño que el cuerpo."""
    doc.styles["Heading 3"].paragraph_format.first_line_indent = Twips(720)
    doc.styles["Heading 4"].paragraph_format.first_line_indent = Twips(1440)
    pie = doc.styles["Caption"]
    pie.font.size = Pt(12)
    pie.font.name = FUENTE


def uniformar_tablas(doc):
    """Las tablas originales estaban a 9 pt; todas pasan a Times New Roman 12
    y sus columnas se reparten de nuevo para que nada se parta a ese tamaño."""
    runs = 0
    for t in doc.tables:
        for fila in t.rows:
            for celda in fila.cells:
                for p in celda.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(TAMANO_TABLA_PT)
                        r.font.name = FUENTE
                        runs += 1
        tiene_combinadas = t._tbl.find(".//" + qn("w:gridSpan")) is not None or t._tbl.find(".//" + qn("w:vMerge")) is not None
        if not tiene_combinadas:
            filas = [[c.text for c in fila.cells] for fila in t.rows]
            _fijar_ancho(t, anchos_columnas(filas))
    return runs


def actualizar_introduccion(doc):
    cambios = {
        "El documento desarrolla únicamente los entregables de la fase de ideación, análisis y diseño.":
            "El documento desarrolla los entregables de las fases de ideación, análisis y diseño, y el avance de "
            "validación técnica y construcción del sistema. El Capítulo I presenta el planteamiento, los objetivos "
            "y los alcances del proyecto. El Capítulo II integra el marco teórico y el estado del arte sobre "
            "pronóstico de ventas minoristas, series de tiempo, aprendizaje automático, demanda intermitente, "
            "métricas y explicabilidad. El Capítulo III define el enfoque metodológico, la población, la muestra y "
            "las técnicas de recolección de datos.",
        "Como complemento de ingeniería se incluye el levantamiento preliminar de requerimientos":
            "Como complemento de ingeniería, la sección 3.5 incluye el levantamiento preliminar de requerimientos "
            "mediante historias de usuario y casos de uso, el diseño arquitectónico lógico y de datos, y tres "
            "wireframes de baja fidelidad. Sobre esa base, el Capítulo IV presenta los resultados de la validación técnica del "
            "prototipo y la comprobación de hipótesis; el Capítulo V, el estudio de factibilidad; y el Capítulo VI, "
            "el desarrollo de la solución: tipo de sistema, requerimientos, diseño, implementación y plan de pruebas.",
    }
    hechos = 0
    for p in doc.paragraphs:
        for inicio, nuevo in cambios.items():
            if p.text.startswith(inicio):
                for r in p.runs[1:]:
                    r._r.getparent().remove(r._r)
                p.runs[0].text = nuevo
                hechos += 1
    return hechos


# ---------------------------------------------------------------------------

def main():
    RESPALDO.mkdir(exist_ok=True)
    for ext in ("docx", "pdf"):
        origen = DOCS / f"Trabajo de Graduacion.{ext}"
        destino = RESPALDO / f"Trabajo de Graduacion - v1 agosto 2026.{ext}"
        if origen.exists() and not destino.exists():
            shutil.copy2(origen, destino)
    if not ORIGINAL.exists():
        sys.exit(f"No existe el original {ORIGINAL}")

    doc = Document(ORIGINAL)
    heading_ref = next(p for p in doc.paragraphs if p.text.strip() == "Referencias" and p.style.name == "Heading 1")
    ins = Insertor(doc, heading_ref._p)

    h1 = doc.styles["Heading 1"]
    rompe = h1.paragraph_format.page_break_before is True
    ajustar_estilos(doc)
    n_runs_tablas = uniformar_tablas(doc)
    n_entregable = integrar_entregable(doc)
    estado = {"mermaid": 0}
    for nombre in CAPITULOS:
        insertar_capitulo(ins, TESIS / nombre, estado, rompe)

    n_intro = actualizar_introduccion(doc)
    n_nuevas, n_total = fusionar_referencias(doc, heading_ref)

    doc.save(SALIDA)
    d2 = Document(SALIDA)
    caps = [p for p in d2.paragraphs if p.style.name == "Caption"]
    print(f"escrito: {SALIDA.name}")
    print(f"  tablas: {len(d2.tables)} | pies de tabla/figura: {len(caps)} | introduccion actualizada: {n_intro} parrafos")
    print(f"  referencias: {n_nuevas} nuevas insertadas, {n_total} en total")
    print(f"  Entregable integrado como seccion 3.5: {n_entregable} titulos | runs de tablas originales a 12 pt: {n_runs_tablas}")
    print(f"  Heading 1 con salto de pagina propio: {rompe}")


if __name__ == "__main__":
    main()
