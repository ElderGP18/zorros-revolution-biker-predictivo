"""Utilidades de bajo nivel para editar el Word de la tesis sin perder formato.

Word reparte el texto de un párrafo en varios runs (revisiones, revisión
ortográfica, campos), de modo que una frase visible puede no existir como
cadena contigua en ningún run. Estas funciones buscan sobre el texto completo
del párrafo y editan los runs implicados conservando su formato.
"""
from __future__ import annotations

import re
from typing import Callable, Iterator, List, Optional

from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

ESTILOS_INDICE = ("toc 1", "toc 2", "toc 3", "table of figures")
RE_NUMERO_TITULO = re.compile(r"^(\d+(?:\.\d+)*) ")


def parrafo(doc, elemento) -> Paragraph:
    return Paragraph(elemento, doc)


def es_parrafo(elemento) -> bool:
    return elemento.tag == qn("w:p")


def nivel_titulo(doc, elemento) -> Optional[int]:
    """Nivel del título (1 a 9) o None si el elemento no es un título."""
    if not es_parrafo(elemento):
        return None
    nombre = parrafo(doc, elemento).style.name
    if nombre.startswith("Heading "):
        return int(nombre.split()[1])
    return None


def titulo(doc, texto: str, nivel: Optional[int] = None) -> Paragraph:
    """El único título cuyo texto es exactamente `texto`."""
    hallados = [p for p in doc.paragraphs
                if p.style.name.startswith("Heading") and p.text.strip() == texto
                and (nivel is None or p.style.name == f"Heading {nivel}")]
    if len(hallados) != 1:
        raise LookupError(f"Se esperaba un título '{texto}' y hay {len(hallados)}")
    return hallados[0]


def parrafo_que_empieza(doc, inicio: str) -> Paragraph:
    hallados = [p for p in doc.paragraphs if p.text.startswith(inicio)]
    if len(hallados) != 1:
        raise LookupError(f"Se esperaba un párrafo que empiece con '{inicio[:50]}' y hay {len(hallados)}")
    return hallados[0]


def seccion(doc, inicio: Paragraph) -> List:
    """El título `inicio` y todo lo que le sigue hasta el siguiente título de
    su mismo nivel o superior. Incluye tablas y figuras."""
    nivel = nivel_titulo(doc, inicio._p)
    elementos = [inicio._p]
    for el in inicio._p.itersiblings():
        if el.tag == qn("w:sectPr"):
            break
        n = nivel_titulo(doc, el)
        if n is not None and n <= nivel:
            break
        elementos.append(el)
    return elementos


def mover_antes(elementos: List, ancla) -> None:
    for el in elementos:
        ancla.addprevious(el)


def eliminar(elementos: List) -> None:
    for el in elementos:
        el.getparent().remove(el)


def parrafos_del_documento(doc) -> Iterator[Paragraph]:
    """Párrafos del cuerpo y de las celdas de tabla, en orden de aparición."""
    for el in doc.element.body.iterchildren():
        if es_parrafo(el):
            yield parrafo(doc, el)
        elif el.tag == qn("w:tbl"):
            for fila in Table(el, doc).rows:
                for celda in fila.cells:
                    yield from celda.paragraphs


def reemplazar_texto(p: Paragraph, nuevo: str) -> None:
    """Sustituye todo el texto conservando el formato del primer run."""
    p.runs[0].text = nuevo
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def cambiar_numero(p: Paragraph, viejo: str, nuevo: str) -> None:
    """Cambia el número inicial de un título (p. ej. '2.1' por '2.2.1')."""
    primero = p.runs[0]
    if not primero.text.startswith(viejo + " "):
        raise ValueError(f"El título '{p.text}' no empieza con '{viejo} '")
    primero.text = nuevo + primero.text[len(viejo):]


def sustituir_regex(p: Paragraph, patron: re.Pattern, funcion: Callable[[re.Match], str],
                    desde: int = 0) -> int:
    """Aplica `patron.sub(funcion)` sobre el texto completo del párrafo aunque
    las coincidencias crucen varios runs. Devuelve cuántas cambiaron."""
    nodos = list(p._p.iter(qn("w:t")))
    if not nodos:
        return 0
    inicios, total = [], 0
    for t in nodos:
        inicios.append(total)
        total += len(t.text or "")
    completo = "".join(t.text or "" for t in nodos)
    cambios = 0
    for m in reversed([m for m in patron.finditer(completo) if m.start() >= desde]):
        nuevo = funcion(m)
        if nuevo != m.group(0):
            _reemplazar_rango(nodos, inicios, m.start(), m.end(), nuevo)
            cambios += 1
    return cambios


def sustituir_texto(p: Paragraph, viejo: str, nuevo: str) -> None:
    if sustituir_regex(p, re.compile(re.escape(viejo)), lambda _m: nuevo) != 1:
        raise ValueError(f"No se encontró una vez '{viejo[:60]}' en '{p.text[:60]}'")


def _reemplazar_rango(nodos, inicios, inicio: int, fin: int, nuevo: str) -> None:
    i = max(k for k, s in enumerate(inicios) if s <= inicio)
    j = max(k for k, s in enumerate(inicios) if s < fin)
    a, b = inicio - inicios[i], fin - inicios[j]
    if i == j:
        texto = nodos[i].text or ""
        _fijar(nodos[i], texto[:a] + nuevo + texto[b:])
        return
    _fijar(nodos[i], (nodos[i].text or "")[:a] + nuevo)
    for k in range(i + 1, j):
        _fijar(nodos[k], "")
    _fijar(nodos[j], (nodos[j].text or "")[b:])


def _fijar(nodo_t, texto: str) -> None:
    nodo_t.text = texto
    nodo_t.set(qn("xml:space"), "preserve")
