"""Alinea la portada y los Capítulos I a III con la estructura oficial de la guía
UMG de Seminario de Tecnología de Información, según la revisión de octubre de 2026.

- Portadas: curso, catedrático y carnés; fecha de entrega; escudo también en la
  segunda portada; el equipo queda en tres integrantes.
- Capítulo I: el planteamiento pasa a 1.2 y cierra con la pregunta de
  investigación; la justificación pasa a 1.3; cinco objetivos específicos;
  1.5 Alcances y Límites con un apartado explícito de límites; 1.6 Hipótesis.
- Capítulo II: 2.1 Estado del Arte, 2.2 Bases Teóricas (los antiguos 2.1 a 2.7
  bajan un nivel) y 2.3 Definición de Términos. La figura de arquitectura sale
  del marco teórico porque corresponde al diseño (Capítulo VI).
- Capítulo III: el procedimiento de recolección pasa a 3.3.6 y se agregan
  3.4 Validación del Instrumento y 3.5 Metodología de Desarrollo del Proyecto.
- El Entregable de Ingeniería deja de ser una sección suelta: sus tres partes
  se trasladan a 6.2.2, 6.3.1 y 6.3.6 (ver `colocar_entregable`).

El contenido nuevo vive en `revision-capitulos-1-a-3.md`, dividido en bloques.
"""
from __future__ import annotations

import copy
import re
from pathlib import Path
from typing import Callable, Dict, List

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from utilidades_docx import (
    RE_NUMERO_TITULO, cambiar_numero, eliminar, mover_antes, nivel_titulo, parrafo,
    parrafo_que_empieza, reemplazar_texto, seccion, sustituir_texto, titulo,
)

RE_BLOQUE = re.compile(r"^<!-- bloque: ([\w-]+) -->$")

CURSO = "CURSO: SEMINARIO DE TECNOLOGÍA DE INFORMACIÓN"
CATEDRATICO = "CATEDRÁTICO: ING. DICSOR HERNÁNDEZ"
CARNES = {
    "GABRIEL ENRIQUE PÉREZ MEZA": "0900-23-14509",
    "LUIS PEDRO ORDOÑEZ CHUTAN": "0900-19-9883",
    "ELDER ESTUARDO GARCIA PACHECO": "0900-24-9106",
}
# Figura en el original de agosto pero ya no integra el equipo (octubre de 2026).
EX_INTEGRANTES = ("BRYAN MANUEL ALBERTO MARTINEZ BARRIOS",)
FECHA_ORIGINAL = "GUATEMALA, AGOSTO 2026"
FECHA = "GUATEMALA, OCTUBRE 2026"
# Posición vertical del escudo en la segunda portada, en EMU, tal como lo
# ubicó el equipo en Word.
DESPLAZAMIENTO_ESCUDO_PORTADILLA = "367665"

CAPITULO_II = "Capítulo II – Marco Teórico"
ENTREGABLE = "Entregable de Ingeniería – Documentación Técnica Preliminar"
# Parte del Entregable -> (número anterior, número nuevo, título nuevo o None).
PARTES_ENTREGABLE = {
    "requerimientos": ("4.1", "6.2.2", "6.2.2 Levantamiento Preliminar de Requerimientos"),
    "diseno": ("4.2", "6.3.1", None),
    "prototipado": ("4.3", "6.3.6", None),
}

Insertar = Callable[[List[str], object], None]


def leer_bloques(ruta: Path) -> Dict[str, List[str]]:
    bloques: Dict[str, List[str]] = {}
    actual = None
    for linea in ruta.read_text(encoding="utf-8").splitlines():
        m = RE_BLOQUE.match(linea)
        if m:
            actual = bloques.setdefault(m.group(1), [])
        elif actual is not None:
            actual.append(linea)
    return bloques


# ---------------------------------------------------------------------------
# Portada
# ---------------------------------------------------------------------------

def completar_portada(doc) -> None:
    """Agrega curso, catedrático y carnés en la primera portada, actualiza la
    fecha y el equipo en ambas portadas y repite el escudo en la segunda. Ocupa
    párrafos vacíos existentes para que la portada siga cabiendo en una página."""
    parrafos = doc.paragraphs
    primer_integrante = next(iter(CARNES))
    i_titulo = next(i for i, p in enumerate(parrafos) if "SISTEMA PREDICTIVO" in p.text)
    i_nombre = next(i for i, p in enumerate(parrafos) if p.text.strip() == primer_integrante)
    vacios = [p for p in parrafos[i_titulo + 1:i_nombre] if not p.text.strip()]
    if len(vacios) < 8:
        raise ValueError(f"La portada tiene {len(vacios)} párrafos vacíos; se necesitan 8")
    formato = parrafos[i_nombre].runs[0]._r.rPr
    for p, texto in ((vacios[2], CURSO), (vacios[3], CATEDRATICO)):
        _escribir_en_portada(p, texto, formato, negrita=True)
    eliminar([p._p for p in vacios[-2:]])
    eliminar([p._p for p in doc.paragraphs if p.text.strip() in EX_INTEGRANTES])
    for nombre, carne in CARNES.items():
        p = next(p for p in doc.paragraphs if p.text.strip() == nombre)
        linea = p.insert_paragraph_before("")
        p._p.addnext(linea._p)
        _escribir_en_portada(linea, f"CARNÉ: {carne}", formato, negrita=False)
    for p in doc.paragraphs:
        if p.text.strip() == FECHA_ORIGINAL:
            reemplazar_texto(p, FECHA)
    _repetir_escudo_en_portadilla(doc)


def _repetir_escudo_en_portadilla(doc) -> None:
    dibujo = next(d for d in doc.element.body.iter(qn("w:drawing")))
    run = copy.deepcopy(dibujo.getparent())
    ancla = run.find(".//" + qn("wp:anchor"))
    for atributo in [a for a in ancla.attrib if a.endswith("}anchorId") or a.endswith("}editId")]:
        del ancla.attrib[atributo]
    ancla.find(qn("wp:positionV")).find(qn("wp:posOffset")).text = DESPLAZAMIENTO_ESCUDO_PORTADILLA
    ids = [int(d.get("id")) for d in doc.element.body.iter(qn("wp:docPr"))]
    run.find(".//" + qn("wp:docPr")).set("id", str(max(ids) + 1))
    titulo_portadilla = [p for p in doc.paragraphs if "SISTEMA PREDICTIVO" in p.text][1]
    titulo_portadilla._p.pPr.addnext(run)


def _escribir_en_portada(p, texto, formato_rpr, negrita: bool) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(texto)
    if formato_rpr is not None:
        run._r.insert(0, copy.deepcopy(formato_rpr))
    run.bold = negrita


# ---------------------------------------------------------------------------
# Capítulos I a III
# ---------------------------------------------------------------------------

def reestructurar_capitulos_1_a_3(doc, bloques: Dict[str, List[str]], insertar: Insertar) -> None:
    _capitulo_1(doc, bloques, insertar)
    _capitulo_2(doc, bloques, insertar)
    _capitulo_3(doc, bloques, insertar)


def _capitulo_1(doc, bloques, insertar) -> None:
    justificacion = titulo(doc, "1.2 Justificación")
    planteamiento = titulo(doc, "1.3 Planteamiento del Problema")
    interrogante = titulo(doc, "1.4 Interrogante de la Investigación")
    mover_antes(seccion(doc, planteamiento), justificacion._p)
    cambiar_numero(planteamiento, "1.3", "1.2")
    cambiar_numero(justificacion, "1.2", "1.3")
    eliminar(seccion(doc, interrogante))
    insertar(bloques["pregunta"], justificacion._p)

    for viejo, nuevo, texto in (("1.5", "1.4", "Objetivos"), ("1.5.1", "1.4.1", "Objetivo General"),
                                ("1.5.2", "1.4.2", "Objetivos Específicos")):
        cambiar_numero(titulo(doc, f"{viejo} {texto}"), viejo, nuevo)
    _reemplazar_objetivos(doc, [l for l in bloques["objetivos"] if l.strip()])

    reemplazar_texto(titulo(doc, "1.6 Alcances"), "1.5 Alcances y Límites")
    for k, texto in enumerate(("Ámbito Geográfico", "Ámbito Tecnológico", "Ámbito Personal",
                               "Ámbito Temporal", "Ámbito Temático"), start=1):
        cambiar_numero(titulo(doc, f"1.6.{k} {texto}"), f"1.6.{k}", f"1.5.{k}")
    tematico = parrafo_que_empieza(doc, "El estudio se concentra en pronóstico de ventas")
    sustituir_texto(tematico, " No aborda recomendadores de productos, segmentación individual de clientes "
                              "ni optimización automática de precios.",
                    " Las exclusiones se declaran de forma explícita en la sección 1.5.6.")
    capitulo_2 = titulo(doc, CAPITULO_II)._p
    insertar(bloques["limites"], capitulo_2)
    insertar(bloques["hipotesis"], capitulo_2)


def _reemplazar_objetivos(doc, nuevos: List[str]) -> None:
    encabezado = titulo(doc, "1.4.2 Objetivos Específicos")
    actuales = []
    for el in encabezado._p.itersiblings():
        if nivel_titulo(doc, el) is not None:
            break
        p = parrafo(doc, el)
        if re.match(r"^\d\. ", p.text):
            actuales.append(p)
    if len(actuales) < len(nuevos):
        raise ValueError("Hay menos objetivos en el documento que en el bloque")
    for p, texto in zip(actuales, nuevos):
        reemplazar_texto(p, texto)
    eliminar([p._p for p in actuales[len(nuevos):]])


def _capitulo_2(doc, bloques, insertar) -> None:
    estado_arte = titulo(doc, "2.8 Estado del Arte")
    primera_base = titulo(doc, "2.1 Gestión de Ventas en el Sector Retail")
    _quitar_figura_del_marco_teorico(doc)
    sustituir_texto(parrafo_que_empieza(doc, "El análisis teórico confirma que la principal"),
                    "El análisis teórico confirma", "La revisión del estado del arte confirma")

    # Las bases teóricas bajan un nivel antes de mover nada: 2.k -> 2.2.k.
    capitulo_3 = titulo(doc, "Capítulo III – Marco Metodológico")
    for el in list(primera_base._p.itersiblings(preceding=False)):
        if el is estado_arte._p or el is capitulo_3._p:
            break
        _bajar_un_nivel(doc, el, "2.", "2.2.")
    _bajar_un_nivel(doc, primera_base._p, "2.", "2.2.")

    mover_antes(seccion(doc, estado_arte), primera_base._p)
    cambiar_numero(estado_arte, "2.8", "2.1")
    cambiar_numero(titulo(doc, "2.8.1 Posicionamiento de la propuesta"), "2.8.1", "2.1.1")
    insertar(bloques["bases-teoricas"], primera_base._p)
    insertar(bloques["definicion-terminos"], capitulo_3._p)


def _bajar_un_nivel(doc, el, prefijo_viejo: str, prefijo_nuevo: str) -> None:
    nivel = nivel_titulo(doc, el)
    if nivel not in (2, 3):
        return
    p = parrafo(doc, el)
    numero = RE_NUMERO_TITULO.match(p.text).group(1)
    cambiar_numero(p, numero, prefijo_nuevo + numero[len(prefijo_viejo):])
    p.style = doc.styles[f"Heading {nivel + 1}"]


def _quitar_figura_del_marco_teorico(doc) -> None:
    """La arquitectura de la solución aparecía dos veces; se conserva solo la
    del diseño (Capítulo VI)."""
    pie = next(p for p in doc.paragraphs if p.style.name == "Caption"
               and p.text.startswith("Figura 1. Arquitectura lógica propuesta"))
    imagen = pie._p.getprevious()
    if imagen.find(".//" + qn("w:drawing")) is None:
        raise ValueError("El párrafo anterior al pie de la Figura 1 no contiene la imagen")
    eliminar([imagen, pie._p])
    sustituir_texto(parrafo_que_empieza(doc, "La arquitectura propuesta separa el servicio de datos"),
                    "La arquitectura propuesta separa",
                    "La arquitectura propuesta para este proyecto, que desarrolla la sección 6.3, separa")


def _capitulo_3(doc, bloques, insertar) -> None:
    reemplazar_texto(titulo(doc, "3.1 Enfoque y Diseño Metodológico"), "3.1 Tipo y Diseño de Investigación")
    insertar(bloques["diseno-comparativo"], titulo(doc, "3.1.1 Unidad de Análisis")._p)
    procedimiento = titulo(doc, "3.4 Procedimiento de Recolección de Datos")
    cambiar_numero(procedimiento, "3.4", "3.3.6")
    procedimiento.style = doc.styles["Heading 3"]
    fin_capitulo = titulo(doc, ENTREGABLE)._p
    insertar(bloques["validacion-instrumento"], fin_capitulo)
    insertar(bloques["metodologia-desarrollo"], fin_capitulo)


# ---------------------------------------------------------------------------
# Entregable de Ingeniería -> Capítulo VI
# ---------------------------------------------------------------------------

def extraer_entregable(doc) -> Dict[str, List]:
    """Separa las tres partes del Entregable, ya renumeradas para el Capítulo
    VI, y elimina su título de primer nivel."""
    encabezado = titulo(doc, ENTREGABLE, nivel=1)
    partes = {}
    for clave, (viejo, nuevo, titulo_nuevo) in PARTES_ENTREGABLE.items():
        inicio = next(p for p in doc.paragraphs
                      if p.style.name == "Heading 2" and p.text.startswith(viejo + " "))
        elementos = seccion(doc, inicio)
        for el in elementos:
            _renumerar_titulo_entregable(doc, el, viejo, nuevo)
        if titulo_nuevo:
            reemplazar_texto(inicio, titulo_nuevo)
        eliminar(elementos)
        partes[clave] = elementos
    eliminar([encabezado._p])
    return partes


def _renumerar_titulo_entregable(doc, el, viejo: str, nuevo: str) -> None:
    nivel = nivel_titulo(doc, el)
    if nivel not in (2, 3):
        return
    p = parrafo(doc, el)
    numero = RE_NUMERO_TITULO.match(p.text).group(1)
    if not numero.startswith(viejo):
        raise ValueError(f"Título inesperado en el Entregable: {p.text}")
    cambiar_numero(p, numero, nuevo + numero[len(viejo):])
    p.style = doc.styles[f"Heading {nivel + 1}"]


def colocar_entregable(partes: Dict[str, List], anclas: Dict[str, object]) -> None:
    """Inserta cada parte en el párrafo marcador que dejó el Markdown del
    Capítulo VI (`<!-- ancla: ... -->`) y elimina el marcador."""
    for clave, elementos in partes.items():
        ancla = anclas[clave]._p
        mover_antes(elementos, ancla)
        eliminar([ancla])
