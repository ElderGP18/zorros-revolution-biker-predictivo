"""Pies de tabla y figura en formato APA y numeración final de tablas y figuras.

- Cada pie queda como «**Tabla N.** *Título*» y la fuente pasa a un párrafo
  propio «*Nota.* …» debajo, en lugar de «Fuente: …» dentro del pie. Así los
  índices de tablas y figuras muestran solo el número y el título.
- Las fuentes (Word de agosto y capítulos en Markdown) citan tablas y figuras
  con su numeración de origen. Tras reordenar el documento, cada pie recibe su
  número final por orden de aparición y todas las menciones «Tabla N» o
  «Figuras N a M» del texto se traducen con ese mismo mapa.
"""
from __future__ import annotations

import re
from typing import Dict, List

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from utilidades_docx import ESTILOS_INDICE, RE_NUMERO_TITULO, parrafos_del_documento, sustituir_regex

FUENTE = "Times New Roman"
TAMANO_INDICE_PT = 12
RE_PIE = re.compile(r"^(Tabla|Figura) (\d+)\.\s*(.*?)\s*(?:Fuente:\s*(.*?))?\s*$", re.S)
RE_REFERENCIA = re.compile(r"\b(Tabla|Figura)(s?)\s+(\d+)((?:\s*(?:,|y|a)\s*\d+)*)")
RE_SECCION = re.compile(r"\bsecci[oó]n(?:es)?\s+(\d+(?:\.\d+)+)(?:\s*(?:,|y|a)\s*(\d+(?:\.\d+)+))?")


def _pies(doc) -> List:
    return [p for p in doc.paragraphs if p.style.name == "Caption" and RE_PIE.match(p.text)]


# ---------------------------------------------------------------------------
# Formato APA
# ---------------------------------------------------------------------------

def formato_apa(doc) -> int:
    notas = 0
    for p in _pies(doc):
        m = RE_PIE.match(p.text)
        titulo_pie, fuente = m.group(3).rstrip(), m.group(4)
        if titulo_pie.endswith("."):
            titulo_pie = titulo_pie[:-1]
        _rehacer_pie(p, titulo_pie)
        if fuente:
            _agregar_nota(p, fuente.strip())
            notas += 1
    return notas


def _rehacer_pie(p, titulo_pie: str) -> None:
    """Conserva «Tabla » y el campo SEQ; reemplaza el resto por el título en cursiva."""
    fin_campo = [f for f in p._p.iter(qn("w:fldChar")) if f.get(qn("w:fldCharType")) == "end"][0]
    run_fin = fin_campo.getparent()
    for el in list(run_fin.itersiblings()):
        if el.tag in (qn("w:r"), qn("w:proofErr"), qn("w:hyperlink")):
            p._p.remove(el)
    for r in p.runs:
        r.bold = True
        r.italic = False
    punto = p.add_run(".")
    punto.bold = True
    cuerpo = p.add_run(" " + titulo_pie)
    cuerpo.italic = True


def _agregar_nota(pie, fuente: str) -> None:
    nota = pie.insert_paragraph_before("", style="Caption")
    pie._p.addnext(nota._p)
    pie.paragraph_format.keep_with_next = True
    nota.alignment = WD_ALIGN_PARAGRAPH.CENTER
    etiqueta = nota.add_run("Nota.")
    etiqueta.italic = True
    nota.add_run(" " + fuente[0].upper() + fuente[1:])


# ---------------------------------------------------------------------------
# Numeración final
# ---------------------------------------------------------------------------

def renumerar(doc) -> Dict[str, Dict[int, int]]:
    mapa: Dict[str, Dict[int, int]] = {"Tabla": {}, "Figura": {}}
    for p in _pies(doc):
        etiqueta, origen = RE_PIE.match(p.text).group(1), int(RE_PIE.match(p.text).group(2))
        if origen in mapa[etiqueta]:
            raise ValueError(f"Dos pies usan el número de origen {etiqueta} {origen}")
        mapa[etiqueta][origen] = len(mapa[etiqueta]) + 1
        _fijar_numero_seq(p, mapa[etiqueta][origen])

    def traducir(m: re.Match) -> str:
        tabla = mapa[m.group(1)]
        texto = m.group(0)
        numeros = [int(n) for n in re.findall(r"\d+", texto)]
        faltan = [n for n in numeros if n not in tabla]
        if faltan:
            raise KeyError(f"'{texto}' cita {m.group(1)} {faltan}, que no existe en el documento")
        return re.sub(r"\d+", lambda n: str(tabla[int(n.group(0))]), texto)

    for p in parrafos_del_documento(doc):
        if p.style.name in ESTILOS_INDICE:
            continue
        # En un pie, la primera coincidencia es su propio número (campo SEQ).
        sustituir_regex(p, RE_REFERENCIA, traducir, desde=1 if p.style.name == "Caption" else 0)
    return mapa


def _fijar_numero_seq(p, numero: int) -> None:
    """Valor visible del campo SEQ; Word lo recalcula igual al actualizar campos."""
    en_resultado = False
    for el in p._p.iter(qn("w:fldChar"), qn("w:t")):
        if el.tag == qn("w:fldChar"):
            tipo = el.get(qn("w:fldCharType"))
            en_resultado = tipo == "separate"
            if tipo == "end":
                return
        elif en_resultado:
            el.text = str(numero)
            return


# ---------------------------------------------------------------------------
# Índices y verificación
# ---------------------------------------------------------------------------

def corregir_estilo_indices(doc) -> None:
    """El estilo de los índices de tablas y figuras traía «mayúsculas» activado,
    y Word guardaba las entradas con cada palabra en mayúscula («Mase Por
    Ventana»). Se quita y se usa la fuente del documento."""
    estilo = doc.styles["table of figures"]
    rpr = estilo.element.get_or_add_rPr()
    for caps in rpr.findall(qn("w:caps")):
        rpr.remove(caps)
    for fuentes in rpr.findall(qn("w:rFonts")):
        rpr.remove(fuentes)
    estilo.font.name = FUENTE
    estilo.font.size = Pt(TAMANO_INDICE_PT)


def verificar(doc) -> List[str]:
    """Problemas que el pipeline no puede resolver solo: secciones citadas que
    no existen, pies con «Fuente:» y títulos con numeración fuera de secuencia."""
    problemas = []
    numeros = []
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            m = RE_NUMERO_TITULO.match(p.text)
            if m:
                numeros.append(m.group(1))
    existentes = set(numeros)
    for anterior, actual in zip(numeros, numeros[1:]):
        if not _sigue(anterior, actual):
            problemas.append(f"numeración: '{actual}' después de '{anterior}'")
    for p in parrafos_del_documento(doc):
        if p.style.name in ESTILOS_INDICE:
            continue
        for m in RE_SECCION.finditer(p.text):
            for citado in filter(None, m.groups()):
                if citado not in existentes:
                    problemas.append(f"sección inexistente {citado}: «{p.text[max(0, m.start() - 40):m.end() + 20]}»")
        if p.style.name == "Caption" and "Fuente:" in p.text:
            problemas.append(f"pie con 'Fuente:': {p.text[:60]}")
    return problemas


def _sigue(anterior: str, actual: str) -> bool:
    a = [int(x) for x in anterior.split(".")]
    b = [int(x) for x in actual.split(".")]
    if b[0] == a[0] + 1:
        return all(x == 1 for x in b[1:])
    if len(b) == len(a) + 1:
        return b[:-1] == a and b[-1] == 1
    if len(b) <= len(a):
        return b[:-1] == a[:len(b) - 1] and b[-1] == a[len(b) - 1] + 1
    return False
